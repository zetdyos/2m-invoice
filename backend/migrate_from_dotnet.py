"""Import vendors (Vendors.xlsx) and confirmations (ACHConfirmations.dat) from the
original .NET Invoicer into the new SQLite database. Run once after install.

Usage:
  python migrate_from_dotnet.py [--vendors PATH] [--confirmations PATH] [--force]
"""
import argparse
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from sqlmodel import Session, select

sys.path.insert(0, str(Path(__file__).resolve().parent))
from app.database import engine, init_db          # noqa: E402
from app.models import Vendor, Confirmation         # noqa: E402

SEP = re.compile(r"[;,\r\n\t ]+")
DEFAULT_ROOT = Path(__file__).resolve().parent.parent.parent  # /home/zet/Invoicer


def norm_invoices(raw: str) -> str:
    parts = [p for p in SEP.split(raw or "") if p.strip()]
    return "\n".join(parts)


def parse_date(s: str):
    try:
        return datetime.fromisoformat(s).date()
    except Exception:
        return datetime.now(timezone.utc).date()


def import_vendors(path: Path, session: Session) -> int:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    count = 0
    for row in ws.iter_rows(values_only=True):
        if row is None or row[0] is None:
            continue
        c0 = str(row[0]).strip()
        if c0.lower() == "id":
            continue
        try:
            vid = int(float(c0))
        except ValueError:
            continue
        name = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
        email = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
        account = str(row[3]).strip() if len(row) > 3 and row[3] is not None else ""
        if not name:
            continue
        session.merge(Vendor(id=vid, name=name, email=email, account=account, active=True))
        count += 1
    wb.close()
    return count


def import_confirmations(path: Path, session: Session) -> int:
    data = json.loads(Path(path).read_text(encoding="utf-8-sig"))
    count = 0
    for r in data:
        session.merge(Confirmation(
            id=int(r["Id"]),
            vendor_id=int(r["VendorId"]),
            confirmation_number=str(r.get("ConfirmationNumber") or ""),
            invoices=norm_invoices(r.get("Invoices") or ""),
            date=parse_date(r.get("Date")),
            amount=float(r.get("Amount") or 0.0),
            status="sent",
            created_by="migration",
        ))
        count += 1
    return count


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--vendors", default=str(DEFAULT_ROOT / "Vendors" / "Vendors.xlsx"))
    ap.add_argument("--confirmations", default=str(DEFAULT_ROOT / "ACHConfirmations.dat"))
    ap.add_argument("--force", action="store_true", help="import even if data already exists")
    args = ap.parse_args()

    init_db()
    with Session(engine) as session:
        existing = session.exec(select(Vendor)).first()
        if existing and not args.force:
            print("Database already has vendors. Use --force to re-import. Aborting.")
            return
        v = import_vendors(Path(args.vendors), session) if Path(args.vendors).exists() else 0
        if not Path(args.vendors).exists():
            print(f"WARN: vendors file not found: {args.vendors}")
        c = import_confirmations(Path(args.confirmations), session) if Path(args.confirmations).exists() else 0
        if not Path(args.confirmations).exists():
            print(f"WARN: confirmations file not found: {args.confirmations}")
        session.commit()
    print(f"Imported {v} vendors and {c} confirmations.")


if __name__ == "__main__":
    main()
